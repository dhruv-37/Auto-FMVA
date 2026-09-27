"""Step 5 — assembling the output. Excel-native path (per the doc: openpyxl/xlsxwriter
builds a workbook with narrative text inserted into designated cells). The other path the
doc mentions (Jinja2 HTML / python-docx / python-pptx polished report) is intentionally
left for a follow-up — this module ships the workbook path fully, since that's the one
that's auditable/re-forecastable, which the doc calls out as the reason to keep it.
"""
from __future__ import annotations
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = {"high": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
             "medium": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
             "low": PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")}


def _excel_safe(value: Any) -> Any:
    """openpyxl can't write pandas' NA sentinels (pd.NA, NaT) or numpy scalar types directly."""
    if value is None or value is pd.NA or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value) if not isinstance(value, (list, dict)) else False:
        return None
    if hasattr(value, "item"):  # numpy scalar -> native python
        try:
            return value.item()
        except (ValueError, TypeError):
            return str(value)
    return value


def _write_df_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(sheet_name[:31])
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in df.itertuples(index=False):
        ws.append([_excel_safe(v) for v in row])
    for i, col in enumerate(df.columns, start=1):
        max_len = df[col].astype(str).str.len().max()
        max_len = int(max_len) if pd.notna(max_len) else 10
        width = max(10, min(28, max_len + 2, len(str(col)) + 4))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _write_narrative_sheet(wb: Workbook, narratives: dict[str, dict]) -> None:
    ws = wb.create_sheet("Narrative")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    row = 1
    for module, parsed in narratives.items():
        ws.cell(row=row, column=1, value=module).font = Font(bold=True, size=13)
        row += 1
        if "title" in parsed:
            ws.cell(row=row, column=1, value="Title").font = Font(bold=True)
            ws.cell(row=row, column=2, value=parsed["title"])
            row += 1
        if "body" in parsed:
            ws.cell(row=row, column=1, value="Body").font = Font(bold=True)
            cell = ws.cell(row=row, column=2, value=parsed["body"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        if "recent_updates" in parsed and parsed["recent_updates"]:
            ws.cell(row=row, column=1, value="Recent Updates").font = Font(bold=True)
            row += 1
            for update in parsed["recent_updates"]:
                ws.cell(row=row, column=2, value=f"• {update}")
                row += 1
        if "flags" in parsed and parsed["flags"]:
            ws.cell(row=row, column=1, value="Flags").font = Font(bold=True)
            row += 1
            for f in parsed["flags"]:
                ws.cell(row=row, column=2, value=json.dumps(f, default=str))
                row += 1
        row += 1  # blank line between modules


def _write_qa_sheet(wb: Workbook, qa_result: dict) -> None:
    ws = wb.create_sheet("QA Review")
    ws.append(["Severity", "Table", "Column", "Year", "Issue"])
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for f in qa_result.get("flags", []):
        r = ws.max_row + 1
        ws.append([f.get("severity"), f.get("table"), f.get("column"), f.get("year"), f.get("issue")])
        fill = FLAG_FILL.get(f.get("severity"))
        if fill:
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = fill
    ws.column_dimensions["E"].width = 90
    for col in "ABCD":
        ws.column_dimensions[col].width = 16
    status_row = ws.max_row + 2
    ws.cell(row=status_row, column=1,
            value=f"Status: {'READY (no high-severity flags)' if qa_result.get('ok') else 'BLOCKED — resolve high-severity flags'}"
            ).font = Font(bold=True, color="B91C1C" if not qa_result.get("ok") else "047857")


def render_excel(out_path: str, tables: dict[str, pd.DataFrame],
                  narratives: dict[str, dict], qa_result: dict) -> str:
    """`tables` — {sheet_name: DataFrame} for every Bucket-A table you want in the workbook
    (financials, ratios, market_data, dcf, peer_comps, monte_carlo_var_summary, ...).
    `narratives` — {module_name: parsed narrative dict} from Step 3/4.
    `qa_result` — output of qa.qa_check(). Writes the workbook and returns its path.
    """
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    for name, df in tables.items():
        if df is not None and not df.empty:
            _write_df_sheet(wb, name, df)
    if narratives:
        _write_narrative_sheet(wb, narratives)
    _write_qa_sheet(wb, qa_result)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
